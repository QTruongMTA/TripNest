from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "2. Nghien cuu va khao sat"
OUT_FILE = OUT_DIR / "01. Phieu khao sat khach du lich.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in [
        ("Heading 1", 20, 18, 6),
        ("Heading 2", 16, 14, 6),
        ("Heading 3", 14, 12, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)


def add_question(doc, text, options):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    for option in options:
        opt = doc.add_paragraph()
        opt.paragraph_format.left_indent = Inches(0.25)
        opt.paragraph_format.space_after = Pt(2)
        r = opt.add_run(f"[ ] {option}")
        r.font.name = "Arial"
        r.font.size = Pt(11)


def add_text_field(doc, label, lines=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(label)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    for _ in range(lines):
        line = doc.add_paragraph()
        line.paragraph_format.space_after = Pt(4)
        r = line.add_run("." * 95)
        r.font.name = "Arial"
        r.font.size = Pt(10)


def add_likert_table(doc):
    doc.add_paragraph()
    caption = doc.add_paragraph()
    caption.paragraph_format.space_after = Pt(4)
    r = caption.add_run("Bảng đánh giá mức độ đồng ý")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(11)

    headers = [
        "Tiêu chí đánh giá",
        "1\nRất không đồng ý",
        "2\nKhông đồng ý",
        "3\nBình thường",
        "4\nĐồng ý",
        "5\nRất đồng ý",
    ]
    rows = [
        "Tôi thường đặt phòng/dịch vụ du lịch qua website hoặc ứng dụng.",
        "Tôi quan tâm đến giao diện dễ sử dụng khi đặt dịch vụ du lịch.",
        "Tôi cần xem hình ảnh, giá, tiện ích và đánh giá trước khi đặt.",
        "Tôi muốn có bộ lọc theo địa điểm, giá, số khách và loại hình lưu trú.",
        "Tôi cần hệ thống xác nhận đặt chỗ rõ ràng và nhanh chóng.",
        "Tôi muốn theo dõi lịch sử đặt chỗ trong tài khoản cá nhân.",
        "Tôi sẵn sàng sử dụng TripNest nếu thông tin minh bạch và thao tác đơn giản.",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    widths = [3860, 1100, 1100, 1100, 1100, 1100]

    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "F2F4F7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(8)

    for label in rows:
        cells = table.add_row().cells
        for i, width in enumerate(widths):
            set_cell_width(cells[i], width)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(label)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        for i in range(1, 6):
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("[ ]")
            run.font.name = "Arial"
            run.font.size = Pt(9)


def build_doc():
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Phiếu khảo sát khách du lịch")
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    r = subtitle.add_run("Phục vụ nghiên cứu nhu cầu sử dụng hệ thống đặt phòng và dịch vụ du lịch TripNest")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(85, 85, 85)

    doc.add_heading("Mục đích khảo sát", level=1)
    doc.add_paragraph(
        "Khảo sát này được thực hiện nhằm tìm hiểu thói quen, nhu cầu và khó khăn của khách du lịch khi tìm kiếm, đặt phòng và sử dụng dịch vụ du lịch trực tuyến. Kết quả khảo sát là cơ sở để nhóm xây dựng và hoàn thiện hệ thống TripNest."
    )

    doc.add_heading("Hướng dẫn trả lời", level=1)
    for item in [
        "Người trả lời đánh dấu [x] vào phương án phù hợp hoặc điền ý kiến vào phần câu hỏi mở.",
        "Thông tin khảo sát chỉ phục vụ mục đích học tập, nghiên cứu và phát triển sản phẩm.",
        "Thời gian trả lời dự kiến: 5-7 phút.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("A. Thông tin chung", level=1)
    add_question(doc, "1. Độ tuổi của bạn:", ["Dưới 18", "18-24", "25-34", "35-44", "Từ 45 trở lên"])
    add_question(doc, "2. Nghề nghiệp hiện tại:", ["Học sinh/Sinh viên", "Nhân viên văn phòng", "Kinh doanh/Tự do", "Khác"])
    add_question(doc, "3. Bạn thường đi du lịch với ai?", ["Một mình", "Gia đình", "Bạn bè", "Đồng nghiệp/nhóm công ty"])
    add_question(doc, "4. Tần suất đi du lịch trong một năm:", ["1 lần", "2-3 lần", "4-5 lần", "Trên 5 lần"])

    doc.add_heading("B. Thói quen đặt dịch vụ du lịch", level=1)
    add_question(doc, "5. Bạn thường tìm kiếm thông tin du lịch qua kênh nào?", ["Google/Website du lịch", "Mạng xã hội", "Bạn bè/người thân giới thiệu", "Ứng dụng đặt phòng", "Công ty du lịch"])
    add_question(doc, "6. Bạn đã từng đặt phòng hoặc dịch vụ du lịch trực tuyến chưa?", ["Đã từng", "Chưa từng"])
    add_question(doc, "7. Yếu tố nào ảnh hưởng nhiều nhất đến quyết định đặt phòng?", ["Giá cả", "Vị trí", "Hình ảnh", "Đánh giá của khách trước", "Tiện ích", "Chính sách hủy/hoàn tiền"])
    add_question(doc, "8. Bạn thường gặp khó khăn gì khi đặt dịch vụ du lịch trực tuyến?", ["Thông tin không rõ ràng", "Giá thay đổi", "Thanh toán phức tạp", "Khó so sánh lựa chọn", "Thiếu đánh giá đáng tin cậy", "Không gặp khó khăn đáng kể"])

    doc.add_heading("C. Đánh giá nhu cầu sử dụng hệ thống", level=1)
    add_likert_table(doc)

    doc.add_heading("D. Mong muốn đối với hệ thống TripNest", level=1)
    add_question(doc, "9. Tính năng nào bạn muốn có trong một website đặt phòng/du lịch?", ["Tìm kiếm và lọc nâng cao", "Xem bản đồ/vị trí", "Đánh giá và bình luận", "Quản lý lịch sử đặt chỗ", "Thông báo xác nhận đặt chỗ", "Ưu đãi/khuyến mãi"])
    add_question(doc, "10. Bạn mong muốn hình thức thanh toán nào?", ["Thanh toán trực tuyến", "Chuyển khoản", "Thanh toán khi nhận dịch vụ", "Cả ba hình thức"])
    add_question(doc, "11. Theo bạn, điều gì làm tăng độ tin cậy của hệ thống?", ["Thông tin minh bạch", "Ảnh thật", "Đánh giá xác thực", "Hỗ trợ khách hàng nhanh", "Chính sách rõ ràng"])

    doc.add_heading("E. Câu hỏi mở", level=1)
    add_text_field(doc, "12. Bạn mong muốn cải thiện điều gì ở các nền tảng đặt phòng/du lịch hiện nay?", lines=3)
    add_text_field(doc, "13. Nếu sử dụng TripNest, bạn kỳ vọng hệ thống hỗ trợ gì tốt hơn?", lines=3)
    add_text_field(doc, "14. Góp ý khác:", lines=3)

    doc.add_heading("Thông tin tổng hợp sau khảo sát", level=1)
    for item in [
        "Số lượng phiếu phát ra: ........................................................",
        "Số lượng phiếu thu về: .........................................................",
        "Số lượng phiếu hợp lệ: .........................................................",
        "Người phụ trách tổng hợp: ......................................................",
        "Ngày tổng hợp: ................................................................",
    ]:
        doc.add_paragraph(item)

    doc.core_properties.title = "Phiếu khảo sát khách du lịch"
    doc.core_properties.subject = "Nghiên cứu và khảo sát nhu cầu người dùng TripNest"
    doc.core_properties.author = "Nhóm thực tập TripNest"
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    build_doc().save(OUT_FILE)


if __name__ == "__main__":
    main()
