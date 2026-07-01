from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "6. Backend"


DOCS = [
    {
        "file": "01. Tong quan Backend.docx",
        "title": "Tong quan Backend",
        "subtitle": "Mo ta vai tro va pham vi cua backend trong du an TripNest",
        "purpose": "Gioi thieu tong quan thanh phan backend, cong nghe su dung va cac nhom chuc nang chinh ma backend dam nhiem.",
        "sections": [
            ("Vai tro backend", [
                "Cung cap API cho frontend va admin portal su dung.",
                "Xu ly nghiep vu dang nhap, phan quyen, tim kiem, dat phong, thanh toan, thong bao va quan ly du lieu.",
                "Ket noi co so du lieu thong qua Prisma ORM va dam bao tinh nhat quan du lieu.",
                "Kiem soat bao mat request thong qua JWT, middleware xac thuc, RBAC, Helmet va rate limit.",
            ]),
            ("Cong nghe su dung", [
                "Node.js va TypeScript dung de xay dung server.",
                "Express dung de khai bao route, middleware va controller.",
                "Prisma ORM dung de thao tac voi PostgreSQL.",
                "Joi dung de validate du lieu dau vao.",
                "bcrypt va jsonwebtoken dung cho xac thuc nguoi dung.",
                "Nodemailer dung cho cac luong gui email/thong bao neu duoc cau hinh.",
            ]),
            ("Cau truc thu muc chinh", [
                "src/routes: dinh nghia endpoint API.",
                "src/controllers: tiep nhan request va goi service xu ly.",
                "src/services: xu ly nghiep vu chinh.",
                "src/middlewares: xac thuc, phan quyen, validate, rate limit va xu ly loi.",
                "prisma: schema, migration va seed data.",
                "uploads: luu tru tep anh tai len trong qua trinh phat trien.",
            ]),
        ],
    },
    {
        "file": "02. Cau truc source code Backend.docx",
        "title": "Cau truc source code Backend",
        "subtitle": "Mo ta cach to chuc ma nguon backend TripNest",
        "purpose": "Giai thich y nghia cac thu muc/file quan trong de thanh vien moi co the doc va bao tri backend nhanh hon.",
        "sections": [
            ("Nguyen tac to chuc", [
                "Tach route, controller, service va middleware de moi lop co trach nhiem ro rang.",
                "Controller chi nen dieu phoi request/response, logic phuc tap nen nam trong service.",
                "Middleware duoc dung lai cho cac nhu cau chung nhu auth, RBAC, validate va error handling.",
                "Prisma schema va migrations la nguon tham chieu chinh cho cau truc database.",
            ]),
            ("Cac nhom file can nam", [
                "index.ts: diem khoi dong server Express.",
                "lib/prisma.ts: khoi tao Prisma client.",
                "validators: schema validate request bang Joi.",
                "utils: ham tien ich nhu JWT, crypto, response va xu ly anh.",
                "generated/prisma: ma sinh tu Prisma, khong nen sua tay neu khong can thiet.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang mo ta thu muc backend",
                "headers": ["Thu muc/File", "Vai tro", "Ghi chu"],
                "rows": [
                    ["src/routes", "Khai bao API endpoint", "Moi module nen co file route rieng"],
                    ["src/controllers", "Nhan request va tra response", "Goi service de xu ly nghiep vu"],
                    ["src/services", "Xu ly logic chinh", "Noi tap trung nghiep vu dat phong, auth, property"],
                    ["src/middlewares", "Xu ly truoc/sau controller", "Auth, RBAC, validate, error, rate limit"],
                    ["prisma/schema.prisma", "Dinh nghia model database", "Can dong bo voi migration"],
                    ["prisma/migrations", "Lich su thay doi database", "Khong sua tuy tien migration da chay"],
                ],
            }
        ],
    },
    {
        "file": "03. Thiet ke kien truc Backend.docx",
        "title": "Thiet ke kien truc Backend",
        "subtitle": "Mo hinh xu ly request va phan lop trong backend",
        "purpose": "Mo ta luong di cua mot request tu frontend den database va cach backend phan tach trach nhiem giua cac lop.",
        "sections": [
            ("Mo hinh xu ly request", [
                "Client gui request den Express server.",
                "Middleware kiem tra CORS, bao mat, rate limit, auth, RBAC va validate du lieu.",
                "Route dieu huong request den controller phu hop.",
                "Controller goi service de xu ly nghiep vu.",
                "Service thao tac database thong qua Prisma client.",
                "Controller tra response chuan ve client hoac chuyen loi sang error middleware.",
            ]),
            ("Uu diem kien truc", [
                "De mo rong them module moi vi route/controller/service duoc tach rieng.",
                "De kiem thu tung lop, dac biet la service va validator.",
                "Giam lap code nhung van giu cau truc de doc voi du an thuc tap nhom.",
            ]),
            ("Luu y khi phat trien", [
                "Khong dat logic database phuc tap truc tiep trong route.",
                "Request dau vao can di qua validator truoc khi goi service.",
                "Tat ca loi nen duoc dua ve error middleware de response thong nhat.",
            ]),
        ],
    },
    {
        "file": "04. Danh sach API Backend.docx",
        "title": "Danh sach API Backend",
        "subtitle": "Tai lieu tong hop endpoint, method va muc dich su dung",
        "purpose": "Lam noi ghi danh sach API de frontend, admin portal va tester tham chieu khi tich hop.",
        "sections": [
            ("Nhom API chinh", [
                "Auth: dang ky, dang nhap, thong tin nguoi dung va token.",
                "Property: danh sach, chi tiet, anh, tien ich va thong tin noi luu tru.",
                "Booking: tao dat phong, xem lich su, cap nhat trang thai va vong doi tai chinh.",
                "Admin/Host/Operator: cac chuc nang quan tri va van hanh.",
                "Review, Notification, Province, Tour, Promotion: cac chuc nang bo tro cho he thong.",
            ]),
            ("Quy uoc response", [
                "Response thanh cong nen co du lieu ro rang va message neu can.",
                "Response loi nen co status code, message va chi tiet validate neu co.",
                "API can xac thuc nen yeu cau Authorization Bearer token.",
            ]),
        ],
        "tables": [
            {
                "title": "Mau bang danh sach API",
                "headers": ["Nhom", "Endpoint", "Method", "Mo ta", "Auth"],
                "rows": [
                    ["Auth", "/api/auth/login", "POST", "Dang nhap nguoi dung", "Khong"],
                    ["Auth", "/api/auth/register", "POST", "Dang ky tai khoan", "Khong"],
                    ["Property", "/api/properties", "GET", "Lay danh sach noi luu tru", "Khong/Tuy chuc nang"],
                    ["Booking", "/api/bookings", "POST", "Tao yeu cau dat phong", "Co"],
                    ["Review", "/api/reviews", "POST", "Tao danh gia sau khi su dung dich vu", "Co"],
                    ["Admin", "/api/admin/...", "GET/POST", "Quan tri he thong", "Co + RBAC"],
                ],
            }
        ],
    },
    {
        "file": "05. Mo ta xu ly nghiep vu Backend.docx",
        "title": "Mo ta xu ly nghiep vu Backend",
        "subtitle": "Tong hop cac luong nghiep vu quan trong trong backend",
        "purpose": "Ghi lai cach backend xu ly cac nghiep vu chinh de phuc vu bao cao, kiem thu va ban giao.",
        "sections": [
            ("Nghiep vu xac thuc", [
                "Nguoi dung gui thong tin dang nhap/dang ky.",
                "Backend validate du lieu, bam mat khau bang bcrypt va tao JWT khi hop le.",
                "Cac API rieng tu su dung middleware auth de kiem tra token.",
            ]),
            ("Nghiep vu dat phong", [
                "Nguoi dung chon noi luu tru, ngay, so khach va gui yeu cau dat phong.",
                "Backend kiem tra du lieu, tinh gia, tao booking va luu trang thai.",
                "Cac buoc thanh toan/xac nhan duoc cap nhat theo trang thai booking.",
            ]),
            ("Nghiep vu van hanh", [
                "Admin/host/operator xu ly phe duyet, cap nhat thong tin va theo doi booking.",
                "He thong ghi nhan thong bao, lich su va cac thong tin lien quan neu co.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang tom tat nghiep vu",
                "headers": ["Nghiep vu", "Dau vao", "Xu ly chinh", "Dau ra"],
                "rows": [
                    ["Dang nhap", "Email, password", "Validate, so khop mat khau, tao token", "Token va thong tin nguoi dung"],
                    ["Dat phong", "Property, ngay, so khach", "Kiem tra hop le, tinh gia, tao booking", "Booking moi"],
                    ["Quan ly property", "Thong tin noi luu tru", "Validate, luu database, xu ly anh", "Property duoc cap nhat"],
                    ["Danh gia", "Noi dung va diem danh gia", "Kiem tra dieu kien, luu review", "Review hien thi tren he thong"],
                ],
            }
        ],
    },
    {
        "file": "06. Ket noi co so du lieu va Prisma.docx",
        "title": "Ket noi co so du lieu va Prisma",
        "subtitle": "Mo ta database, Prisma schema, migration va seed data",
        "purpose": "Giai thich cach backend ket noi PostgreSQL va quan ly cau truc du lieu bang Prisma.",
        "sections": [
            ("Thanh phan lien quan", [
                "DATABASE_URL trong file moi truong dung de ket noi PostgreSQL.",
                "prisma/schema.prisma dinh nghia model, quan he va enum.",
                "prisma/migrations luu cac lan thay doi cau truc database.",
                "prisma/seed.ts dung de nap du lieu mau cho qua trinh demo/kiem thu.",
            ]),
            ("Quy trinh lam viec voi database", [
                "Cap nhat schema Prisma khi can thay doi model.",
                "Tao/chay migration de dong bo database.",
                "Generate Prisma client neu schema thay doi.",
                "Seed du lieu mau khi can khoi tao moi truong demo.",
            ]),
            ("Luu y", [
                "Khong commit thong tin ket noi that trong .env.",
                "Can sao luu du lieu truoc khi chay migration tren moi truong that.",
                "Migration can co ten ro rang theo y nghia thay doi.",
            ]),
        ],
        "tables": [
            {
                "title": "Lenh tham khao",
                "headers": ["Lenh", "Muc dich", "Ghi chu"],
                "rows": [
                    ["npm run seed", "Nap du lieu mau", "Chay trong thu muc backend"],
                    ["npm run migrate:deploy", "Ap dung migration", "Dung khi trien khai"],
                    ["npx prisma generate", "Sinh Prisma client", "Chay sau khi doi schema"],
                    ["npx prisma studio", "Mo giao dien xem database", "Dung khi debug/dev"],
                ],
            }
        ],
    },
    {
        "file": "07. Bao mat va phan quyen Backend.docx",
        "title": "Bao mat va phan quyen Backend",
        "subtitle": "Cac co che bao ve API va du lieu nguoi dung",
        "purpose": "Tong hop cac bien phap bao mat da/nen ap dung trong backend TripNest.",
        "sections": [
            ("Cac co che bao mat", [
                "JWT dung de xac thuc request cua nguoi dung da dang nhap.",
                "bcrypt dung de bam mat khau, khong luu mat khau dang plain text.",
                "RBAC middleware dung de gioi han API theo vai tro nguoi dung.",
                "Joi validator dung de chan du lieu sai dinh dang truoc khi xu ly nghiep vu.",
                "Helmet giup bo sung HTTP security headers.",
                "Rate limit giup han che request lap lai qua nhieu trong thoi gian ngan.",
            ]),
            ("Nguyen tac can tuan thu", [
                "Khong dua secret, JWT key, DATABASE_URL that vao tai lieu cong khai.",
                "API quan tri phai co auth va kiem tra role.",
                "Thong bao loi khong nen lam lo thong tin noi bo nhu stack trace tren moi truong production.",
                "File upload can gioi han loai file, kich thuoc va duong dan luu tru.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang kiem tra bao mat",
                "headers": ["Hang muc", "Trang thai", "Ghi chu"],
                "rows": [
                    ["Mat khau duoc hash bang bcrypt", "Can kiem tra", ""],
                    ["API rieng tu yeu cau JWT", "Can kiem tra", ""],
                    ["API theo role co RBAC", "Can kiem tra", ""],
                    ["Request dau vao duoc validate", "Can kiem tra", ""],
                    ["Khong commit file .env that", "Can kiem tra", ""],
                ],
            }
        ],
    },
    {
        "file": "08. Huong dan cai dat va chay Backend.docx",
        "title": "Huong dan cai dat va chay Backend",
        "subtitle": "Cac buoc khoi chay backend trong moi truong phat trien",
        "purpose": "Huong dan thanh vien nhom cai dat dependency, cau hinh bien moi truong, database va chay server backend.",
        "sections": [
            ("Yeu cau moi truong", [
                "Node.js va npm.",
                "PostgreSQL hoac database tuong thich voi cau hinh Prisma.",
                "File .env duoc tao tu .env.example va cap nhat dung thong tin ket noi.",
            ]),
            ("Cac buoc cai dat", [
                "Mo terminal tai thu muc backend.",
                "Chay npm install neu chua co node_modules.",
                "Cau hinh DATABASE_URL va cac bien moi truong can thiet.",
                "Chay migration/seed neu can khoi tao database.",
                "Chay npm run dev de khoi dong server phat trien.",
            ]),
            ("Lenh npm trong package", [
                "npm run dev: chay server bang nodemon va tsx.",
                "npm run build: bien dich TypeScript sang thu muc dist.",
                "npm start: chay ban build trong dist.",
                "npm run seed: nap du lieu mau.",
                "npm run migrate:deploy: ap dung migration khi trien khai.",
            ]),
        ],
    },
    {
        "file": "09. Bao cao hoan thien Backend.docx",
        "title": "Bao cao hoan thien Backend",
        "subtitle": "Tong ket ket qua phat trien va ban giao phan backend",
        "purpose": "Lam tai lieu tong ket cac chuc nang backend da hoan thanh, nhung han che con lai va huong phat trien tiep theo.",
        "sections": [
            ("Noi dung can bao cao", [
                "Danh sach module backend da hoan thanh.",
                "Danh sach API da tich hop voi frontend/admin portal.",
                "Cac middleware va co che bao mat da ap dung.",
                "Tinh trang database, migration va seed data.",
                "Cac loi da sua va loi con ton neu co.",
            ]),
            ("Han che va huong phat trien", [
                "Bo sung test tu dong cho service va API.",
                "Hoan thien tai lieu API chi tiet theo request/response.",
                "Toi uu log, monitoring va xu ly loi production.",
                "Kiem tra lai bao mat file upload, rate limit va phan quyen theo vai tro.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang tong ket module backend",
                "headers": ["Module", "Chuc nang chinh", "Trang thai", "Ghi chu"],
                "rows": [
                    ["Auth", "Dang ky, dang nhap, JWT", "Can cap nhat", ""],
                    ["Property", "Quan ly noi luu tru va anh", "Can cap nhat", ""],
                    ["Booking", "Dat phong va vong doi booking", "Can cap nhat", ""],
                    ["Admin/Host/Operator", "Quan tri va van hanh", "Can cap nhat", ""],
                    ["Notification/Email", "Thong bao va email", "Can cap nhat", ""],
                ],
            }
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in [
        ("Heading 1", 20, 20, 6),
        ("Heading 2", 16, 18, 6),
        ("Heading 3", 14, 16, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)


def add_table(doc, spec):
    doc.add_paragraph()
    caption = doc.add_paragraph()
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(spec["title"])
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=len(spec["headers"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    widths = [9360 // len(spec["headers"])] * len(spec["headers"])
    for i, text in enumerate(spec["headers"]):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i])
        set_cell_shading(cell, "F2F4F7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)

    for row in spec["rows"]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)


def build_doc(spec):
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run(spec["title"])
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sub_run = subtitle.add_run(spec["subtitle"])
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(85, 85, 85)

    doc.add_heading("Muc dich", level=1)
    doc.add_paragraph(spec["purpose"])

    for heading, items in spec.get("sections", []):
        doc.add_heading(heading, level=1)
        for item in items:
            add_bullet(doc, item)

    for table in spec.get("tables", []):
        add_table(doc, table)

    doc.add_heading("Checklist hoan thanh", level=1)
    for item in [
        "Da cap nhat thong tin theo source code hien tai.",
        "Da bo sung minh chung nhu endpoint, anh chup API, log hoac ket qua chay thu neu co.",
        "Da thong nhat noi dung voi thanh vien phu trach backend.",
    ]:
        add_bullet(doc, item)

    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = "Tai lieu backend du an TripNest"
    doc.core_properties.author = "Nhom thuc tap TripNest"
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    for spec in DOCS:
        build_doc(spec).save(OUT_DIR / spec["file"])


if __name__ == "__main__":
    main()
