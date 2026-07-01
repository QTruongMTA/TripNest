from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "8. Kiem thu"


DOCS = [
    {
        "file": "01. Tong quan thu muc kiem thu.docx",
        "title": "Tong quan thu muc kiem thu",
        "subtitle": "Mo ta cac tai lieu can co trong folder 8. Kiem thu",
        "purpose": "Xac dinh bo tai lieu phuc vu qua trinh lap ke hoach, thuc hien, ghi nhan loi va tong ket ket qua kiem thu cho du an TripNest.",
        "sections": [
            ("Cau truc de xuat", [
                "01. Tong quan thu muc kiem thu.docx",
                "02. Ke hoach kiem thu.docx",
                "03. Chien luoc va pham vi kiem thu.docx",
                "04. Test case Frontend.docx",
                "05. Test case Backend API.docx",
                "06. Test case Tich hop he thong.docx",
                "07. Bien ban loi va theo doi sua loi.docx",
                "08. Bao cao ket qua kiem thu.docx",
                "09. Bien ban nghiem thu kiem thu.docx",
            ]),
            ("Nguyen tac sap xep", [
                "Dat ten file theo so thu tu de de theo doi tien do.",
                "Moi tai lieu nen co nguoi phu trach, ngay cap nhat va trang thai hoan thanh.",
                "Anh chup man hinh loi, log API va minh chung test nen luu kem trong thu muc con Minh chung kiem thu.",
            ]),
        ],
    },
    {
        "file": "02. Ke hoach kiem thu.docx",
        "title": "Ke hoach kiem thu",
        "subtitle": "Test Plan cho du an TripNest",
        "purpose": "Trinh bay muc tieu, nguon luc, lich trinh va cach thuc quan ly hoat dong kiem thu.",
        "sections": [
            ("Muc tieu kiem thu", [
                "Dam bao cac chuc nang chinh cua he thong hoat dong dung theo dac ta.",
                "Phat hien loi giao dien, nghiep vu, ket noi API va luu tru du lieu.",
                "Danh gia muc do san sang cua san pham truoc khi ban giao.",
            ]),
            ("Pham vi tong quat", [
                "Kiem thu giao dien nguoi dung va luong thao tac tren frontend.",
                "Kiem thu API, xu ly du lieu va rang buoc nghiep vu o backend.",
                "Kiem thu tich hop giua frontend, backend va co so du lieu.",
            ]),
            ("Nhan su va vai tro", [
                "Truong nhom kiem thu: lap ke hoach, phan cong va tong hop bao cao.",
                "Tester Frontend: kiem tra man hinh, form, dieu huong va responsive.",
                "Tester Backend: kiem tra API, request/response, validate va ma loi.",
                "Lap trinh vien: tiep nhan loi, sua loi va cap nhat trang thai.",
            ]),
        ],
        "tables": [
            {
                "title": "Lich trinh kiem thu",
                "headers": ["Giai doan", "Noi dung", "Nguoi phu trach", "Trang thai"],
                "rows": [
                    ["Chuan bi", "Lap test case, chuan bi tai khoan va du lieu mau", "", "Chua thuc hien"],
                    ["Thuc hien", "Chay test case, ghi nhan ket qua va loi phat sinh", "", "Chua thuc hien"],
                    ["Sua loi", "Lap trinh vien xu ly loi va tester kiem tra lai", "", "Chua thuc hien"],
                    ["Tong ket", "Lap bao cao ket qua va bien ban nghiem thu", "", "Chua thuc hien"],
                ],
            }
        ],
    },
    {
        "file": "03. Chien luoc va pham vi kiem thu.docx",
        "title": "Chien luoc va pham vi kiem thu",
        "subtitle": "Dinh huong phuong phap test cho TripNest",
        "purpose": "Mo ta cac loai kiem thu se ap dung, tieu chi vao/ra va nhung noi dung khong nam trong pham vi.",
        "sections": [
            ("Loai kiem thu ap dung", [
                "Kiem thu chuc nang: xac minh tinh dung dan cua tung chuc nang.",
                "Kiem thu giao dien: bo cuc, mau sac, hien thi du lieu, thong bao loi.",
                "Kiem thu API: endpoint, method, status code, validate request va response.",
                "Kiem thu tich hop: luong du lieu giua frontend, backend va database.",
                "Kiem thu hoi quy: dam bao sua loi khong lam hong chuc nang da dung.",
            ]),
            ("Tieu chi bat dau", [
                "Da co ban build co the chay duoc.",
                "Da co danh sach chuc nang va dac ta nghiep vu can kiem tra.",
                "Da chuan bi tai khoan test, du lieu mau va moi truong kiem thu.",
            ]),
            ("Tieu chi ket thuc", [
                "Cac test case muc do quan trong cao da duoc chay.",
                "Loi nghiem trong va loi anh huong nghiep vu chinh da duoc sua.",
                "Bao cao ket qua kiem thu va bien ban nghiem thu da duoc cap nhat.",
            ]),
        ],
    },
    {
        "file": "04. Test case Frontend.docx",
        "title": "Test case Frontend",
        "subtitle": "Danh sach kich ban kiem thu giao dien nguoi dung",
        "purpose": "Ghi lai cac test case danh cho man hinh, form, dieu huong, responsive va xu ly thong bao tren frontend.",
        "sections": [
            ("Pham vi frontend", [
                "Kiem tra trang chu, dang nhap, dang ky, tim kiem, chi tiet phong/dich vu va thao tac dat phong.",
                "Kiem tra form bat buoc, thong bao loi, thong bao thanh cong va dieu huong sau thao tac.",
                "Kiem tra hien thi tren desktop va mobile.",
            ]),
        ],
        "tables": [
            {
                "title": "Mau danh sach test case frontend",
                "headers": ["Ma TC", "Man hinh", "Buoc thuc hien", "Ket qua mong doi", "Ket qua", "Trang thai"],
                "rows": [
                    ["FE-01", "Dang nhap", "Nhap email va mat khau hop le, bam Dang nhap", "Dang nhap thanh cong va chuyen vao trang chinh", "", "Chua test"],
                    ["FE-02", "Dang nhap", "De trong email hoac mat khau", "He thong hien thong bao bat buoc nhap", "", "Chua test"],
                    ["FE-03", "Tim kiem", "Nhap tu khoa, ngay va so khach", "Danh sach ket qua phu hop duoc hien thi", "", "Chua test"],
                    ["FE-04", "Chi tiet", "Mo chi tiet mot phong/dich vu", "Thong tin, gia, hinh anh va nut dat phong hien thi dung", "", "Chua test"],
                    ["FE-05", "Responsive", "Mo giao dien tren kich thuoc mobile", "Bo cuc khong vo, nut bam va chu de doc", "", "Chua test"],
                ],
            }
        ],
    },
    {
        "file": "05. Test case Backend API.docx",
        "title": "Test case Backend API",
        "subtitle": "Kich ban kiem thu API va xu ly nghiep vu backend",
        "purpose": "Kiem tra endpoint, phuong thuc, tham so dau vao, ket qua tra ve va xu ly loi cua backend.",
        "sections": [
            ("Pham vi backend API", [
                "Kiem tra API dang nhap, dang ky, lay danh sach, tim kiem, tao yeu cau dat phong va quan ly du lieu.",
                "Kiem tra request hop le, request thieu truong, request sai dinh dang va request khong co quyen.",
                "Kiem tra status code, message loi va cau truc JSON tra ve.",
            ]),
        ],
        "tables": [
            {
                "title": "Mau danh sach test case backend API",
                "headers": ["Ma TC", "API", "Method", "Du lieu dau vao", "Ket qua mong doi", "Trang thai"],
                "rows": [
                    ["BE-01", "/api/auth/login", "POST", "Email, password hop le", "Tra ve token/thong tin nguoi dung", "Chua test"],
                    ["BE-02", "/api/auth/login", "POST", "Sai password", "Tra ve loi xac thuc phu hop", "Chua test"],
                    ["BE-03", "/api/bookings", "POST", "Du lieu dat phong hop le", "Tao booking thanh cong", "Chua test"],
                    ["BE-04", "/api/bookings", "POST", "Thieu truong bat buoc", "Tra ve loi validate", "Chua test"],
                    ["BE-05", "/api/users/me", "GET", "Khong gui token", "Tra ve loi chua xac thuc", "Chua test"],
                ],
            }
        ],
    },
    {
        "file": "06. Test case Tich hop he thong.docx",
        "title": "Test case Tich hop he thong",
        "subtitle": "Kiem thu luong nghiep vu giua frontend, backend va database",
        "purpose": "Xac minh cac thanh phan cua he thong hoat dong dong bo khi nguoi dung thuc hien mot quy trinh hoan chinh.",
        "sections": [
            ("Luong tich hop can kiem tra", [
                "Nguoi dung dang ky tai khoan moi va dang nhap bang tai khoan do.",
                "Nguoi dung tim kiem, xem chi tiet va gui yeu cau dat phong.",
                "Quan tri vien xem danh sach dat phong va cap nhat trang thai.",
                "Frontend hien thi dung du lieu sau khi backend thay doi.",
            ]),
        ],
        "tables": [
            {
                "title": "Mau test case tich hop",
                "headers": ["Ma TC", "Luong kiem thu", "Du lieu test", "Ket qua mong doi", "Minh chung", "Trang thai"],
                "rows": [
                    ["INT-01", "Dang ky -> Dang nhap", "Tai khoan moi", "Tai khoan duoc tao va dang nhap thanh cong", "", "Chua test"],
                    ["INT-02", "Tim kiem -> Xem chi tiet", "Tu khoa va ngay hop le", "Du lieu hien thi thong nhat giua danh sach va chi tiet", "", "Chua test"],
                    ["INT-03", "Dat phong -> Luu database", "Thong tin booking hop le", "Booking duoc tao va co the truy van lai", "", "Chua test"],
                    ["INT-04", "Admin cap nhat trang thai", "Booking dang cho xu ly", "Nguoi dung thay trang thai moi sau khi tai lai", "", "Chua test"],
                ],
            }
        ],
    },
    {
        "file": "07. Bien ban loi va theo doi sua loi.docx",
        "title": "Bien ban loi va theo doi sua loi",
        "subtitle": "Danh sach bug, muc do anh huong va trang thai xu ly",
        "purpose": "Ghi nhan loi phat hien trong qua trinh kiem thu, phan cong nguoi sua va theo doi ket qua retest.",
        "sections": [
            ("Quy uoc muc do loi", [
                "Critical: loi lam he thong khong the tiep tuc su dung hoac mat du lieu quan trong.",
                "High: loi anh huong truc tiep den nghiep vu chinh.",
                "Medium: loi anh huong mot phan chuc nang nhung co cach xu ly tam.",
                "Low: loi hien thi, noi dung, trai nghiem nguoi dung hoac cai tien nho.",
            ]),
        ],
        "tables": [
            {
                "title": "Mau bang theo doi loi",
                "headers": ["Ma loi", "Mo ta loi", "Muc do", "Nguoi sua", "Trang thai", "Ket qua retest"],
                "rows": [
                    ["BUG-01", "", "High", "", "Open", ""],
                    ["BUG-02", "", "Medium", "", "Open", ""],
                    ["BUG-03", "", "Low", "", "Open", ""],
                    ["BUG-04", "", "Critical", "", "Open", ""],
                ],
            }
        ],
    },
    {
        "file": "08. Bao cao ket qua kiem thu.docx",
        "title": "Bao cao ket qua kiem thu",
        "subtitle": "Tong hop ket qua test va danh gia chat luong san pham",
        "purpose": "Tong ket so luong test case da chay, ty le dat/khong dat, danh sach loi con ton va kien nghi truoc khi ban giao.",
        "sections": [
            ("Noi dung can tong hop", [
                "Tong so test case, so test case dat, khong dat va chua thuc hien.",
                "Danh sach loi theo muc do Critical, High, Medium, Low.",
                "Nhan xet ve tinh on dinh cua frontend, backend va cac luong tich hop.",
                "Kien nghi: cho phep ban giao, ban giao co dieu kien hoac can sua them.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang tong hop ket qua",
                "headers": ["Hang muc", "Tong TC", "Dat", "Khong dat", "Chua test", "Ghi chu"],
                "rows": [
                    ["Frontend", "", "", "", "", ""],
                    ["Backend API", "", "", "", "", ""],
                    ["Tich hop he thong", "", "", "", "", ""],
                    ["Hoi quy sau sua loi", "", "", "", "", ""],
                ],
            }
        ],
    },
    {
        "file": "09. Bien ban nghiem thu kiem thu.docx",
        "title": "Bien ban nghiem thu kiem thu",
        "subtitle": "Xac nhan ket qua kiem thu va dieu kien ban giao",
        "purpose": "Lam can cu xac nhan nhom da hoan thanh hoat dong kiem thu theo pham vi da thong nhat.",
        "sections": [
            ("Thong tin nghiem thu", [
                "Ten du an: TripNest.",
                "Hang muc nghiem thu: Kiem thu chuc nang, giao dien, API va tich hop.",
                "Thoi gian nghiem thu: dien ngay bat dau va ngay ket thuc.",
                "Thanh phan tham gia: dai dien nhom phat trien, nhom kiem thu va giang vien/nguoi huong dan neu co.",
            ]),
            ("Ket luan de xuat", [
                "Dat yeu cau: cac chuc nang chinh hoat dong dung, khong con loi nghiem trong.",
                "Dat co dieu kien: con loi nho nhung khong anh huong nghiep vu chinh.",
                "Chua dat: con loi nghiem trong can sua truoc khi ban giao.",
            ]),
        ],
        "tables": [
            {
                "title": "Bang ky xac nhan",
                "headers": ["Vai tro", "Ho ten", "Y kien", "Chu ky"],
                "rows": [
                    ["Dai dien nhom phat trien", "", "", ""],
                    ["Dai dien nhom kiem thu", "", "", ""],
                    ["Nguoi huong dan/Giang vien", "", "", ""],
                ],
            }
        ],
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(11)


def add_table(doc, spec):
    doc.add_paragraph()
    caption = doc.add_paragraph()
    caption.paragraph_format.space_after = Pt(4)
    run = caption.add_run(spec["title"])
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)

    table = doc.add_table(rows=1, cols=len(spec["headers"]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)

    widths = [9360 // len(spec["headers"])] * len(spec["headers"])
    hdr = table.rows[0].cells
    for i, text in enumerate(spec["headers"]):
        set_cell_width(hdr[i], widths[i])
        set_cell_shading(hdr[i], "F2F4F7")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(9)

    for row in spec["rows"]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)


def configure_doc(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, before, after in [
        ("Heading 1", 20, 20, 6),
        ("Heading 2", 16, 18, 6),
        ("Heading 3", 14, 16, 4),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def build_doc(spec):
    doc = Document()
    configure_doc(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run(spec["title"])
    run.font.name = "Arial"
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    sub_run = subtitle.add_run(spec["subtitle"])
    sub_run.font.name = "Arial"
    sub_run.font.size = Pt(11)
    sub_run.font.color.rgb = RGBColor(85, 85, 85)

    doc.add_heading("Muc dich", level=1)
    doc.add_paragraph(spec["purpose"])

    for heading, items in spec.get("sections", []):
        doc.add_heading(heading, level=1)
        for item in items:
            add_bullet(doc, item)

    for table in spec.get("tables", []):
        add_table(doc, table)

    doc.add_heading("Checklist hoan thanh", level=1)
    for item in [
        "Da cap nhat nguoi phu trach va ngay thuc hien.",
        "Da bo sung minh chung neu co anh chup man hinh, log API hoac file ket qua.",
        "Da thong nhat trang thai cuoi cung voi cac thanh vien lien quan.",
    ]:
        add_bullet(doc, item)

    doc.core_properties.title = spec["title"]
    doc.core_properties.subject = "Tai lieu kiem thu du an TripNest"
    doc.core_properties.author = "Nhom thuc tap TripNest"
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    for spec in DOCS:
        doc = build_doc(spec)
        doc.save(OUT_DIR / spec["file"])


if __name__ == "__main__":
    main()
