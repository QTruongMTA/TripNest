from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = r"C:\Webbooking\TripNest\TripNest_Dac_ta_He_thong_v1.1.docx"


def set_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def style_paragraph(paragraph, before=0, after=6, line=1.1):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    if level == 1:
        r = p.add_run(text)
        set_font(r, size=16, bold=True, color="2E74B5")
        style_paragraph(p, before=12, after=6, line=1.0)
    elif level == 2:
        r = p.add_run(text)
        set_font(r, size=13, bold=True, color="2E74B5")
        style_paragraph(p, before=10, after=4, line=1.0)
    else:
        r = p.add_run(text)
        set_font(r, size=12, bold=True, color="1F4D78")
        style_paragraph(p, before=8, after=3, line=1.0)


def add_para(doc, text):
    p = doc.add_paragraph()
    style_paragraph(p, after=5, line=1.12)
    r = p.add_run(text)
    set_font(r)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    style_paragraph(p, after=4, line=1.12)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(text)
    set_font(r)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    style_paragraph(p, after=4, line=1.12)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    r = p.add_run(text)
    set_font(r)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_font(r, size=10, bold=True)
        set_cell_shading(hdr[i], "F2F4F7")
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        hdr[i].width = widths[i]
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if isinstance(val, list):
                for idx, part in enumerate(val):
                    r = p.add_run(part)
                    set_font(r, size=10)
                    if idx < len(val) - 1:
                        r.add_break()
            else:
                r = p.add_run(str(val))
                set_font(r, size=10)
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].width = widths[i]
    return table


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.45)
sec.footer_distance = Inches(0.45)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
styles["Normal"].font.size = Pt(11)
styles["Normal"].paragraph_format.space_after = Pt(6)
for idx, (size, color) in enumerate([(16, "2E74B5"), (13, "2E74B5"), (12, "1F4D78")], start=1):
    st = styles[f"Heading {idx}"]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(3)
r = title.add_run("TÀI LIỆU ĐẶC TẢ HỆ THỐNG")
set_font(r, size=24, bold=True, color="000000")

sub = doc.add_paragraph()
sub.paragraph_format.space_after = Pt(10)
r = sub.add_run("TripNest - Hệ thống đặt phòng lưu trú trực tuyến")
set_font(r, size=12, color="444444")

add_table(
    doc,
    ["Thông tin", "Chi tiết"],
    [
        ["Tên hệ thống", "TripNest"],
        ["Phiên bản tài liệu", "v1.1 (cập nhật theo hệ thống hiện tại)"],
        ["Ngày cập nhật", "10/06/2026"],
        ["Phạm vi", "Frontend khách, Admin Portal, Backend API, PostgreSQL/Prisma"],
        ["Mục đích", "Đặc tả phản ánh đúng chức năng đã triển khai và các phần còn lại đang để placeholder"],
    ],
    [Inches(1.8), Inches(4.7)],
)

add_heading(doc, "1. Mục tiêu tài liệu", 1)
add_para(doc, "Tài liệu này thay thế phần mô tả cũ chưa còn khớp với mã nguồn. Nội dung đã được cập nhật để phản ánh đúng hệ thống TripNest hiện tại: vai trò người dùng, module, API, mô hình dữ liệu và các phần chưa triển khai hoàn chỉnh.")
add_para(doc, "Tài liệu tách rõ ba nhóm trạng thái: đã triển khai, đang có ở mức một phần/placeholder, và định hướng tương lai.")

add_heading(doc, "2. Kiến trúc hiện tại", 1)
add_table(
    doc,
    ["Tầng", "Công nghệ", "Vai trò"],
    [
        ["Frontend khách", "Next.js 14", "Tìm kiếm, chi tiết cơ sở/tour, đặt chỗ, hồ sơ host"],
        ["Admin Portal", "Next.js 14", "Quản trị, operator, user, listing, audit log"],
        ["Backend API", "Node.js + Express.js", "REST API, RBAC, nghiệp vụ"],
        ["Database", "PostgreSQL + Prisma", "User, Property, Tour, Booking, Payment, Review, Province, Operator..."],
        ["Lưu file", "Upload local", "Ảnh cơ sở lưu trú trong môi trường dev"],
        ["Bản đồ", "OpenStreetMap / Nominatim", "Dùng ở form đăng cơ sở mới"],
    ],
    [Inches(1.3), Inches(1.8), Inches(3.4)],
)

add_heading(doc, "3. Vai trò người dùng", 1)
add_table(
    doc,
    ["Role", "Tên gọi", "Phạm vi chính"],
    [
        ["GUEST", "Khách du lịch", "Tìm kiếm, xem chi tiết, đặt chỗ, theo dõi booking, hồ sơ cá nhân"],
        ["HOST", "Chủ cơ sở", "Đăng cơ sở, cập nhật hồ sơ, quản lý booking và xem doanh thu"],
        ["OPERATOR_PROVINCE", "Operator tỉnh", "Quản lý listing theo tỉnh, duyệt host/listing, xử lý tranh chấp, giao việc"],
        ["OPERATOR_SUB", "Operator thực địa", "Thực hiện task kiểm tra/điều tra theo phân công"],
        ["ADMIN", "Quản trị viên", "Quản lý hệ thống, tỉnh/thành, operator, listing, user, audit log"],
    ],
    [Inches(1.2), Inches(1.5), Inches(3.8)],
)

add_heading(doc, "4. Chức năng đã phản ánh đúng hệ thống", 1)
for item in [
    "Đăng ký / đăng nhập / đăng xuất / xem hồ sơ hiện tại.",
    "Cập nhật hồ sơ cá nhân: tên hiển thị, ảnh đại diện, điện thoại, ngày sinh, quốc tịch, địa chỉ.",
    "Tìm kiếm và xem chi tiết cơ sở lưu trú.",
    "Danh sách tour và chi tiết tour.",
    "Tạo booking cho property, có kiểm tra trùng lịch và trạng thái bị khóa.",
    "Host đăng cơ sở mới với thông tin chi tiết: địa chỉ, tiện nghi, ngôn ngữ, nội quy, ảnh, pháp lý, ownerAlias.",
    "Host quản lý danh sách cơ sở của mình, lịch booking và doanh thu.",
    "Admin quản lý user, listing, booking, payments, promotions, commissions, reviews và audit logs.",
    "Admin tạo operator, phân công tỉnh cho operator tỉnh.",
    "Operator xem dashboard, listing theo tỉnh, host approvals, task, dispute và sub-operators.",
    "Hệ thống có notifications và audit log ở tầng dữ liệu.",
]:
    add_bullet(doc, item)

add_heading(doc, "5. API cần cập nhật trong đặc tả", 1)
add_table(
    doc,
    ["Module", "Endpoint nhóm", "Ghi chú"],
    [
        ["Auth", "/api/v1/auth/register, /login, /logout, /me, /me/profile, /me/avatar, /me/password", "Không có refresh token / OAuth / 2FA thực tế"],
        ["Properties / Tours", "/api/v1/properties, /api/v1/tours", "Tìm kiếm, chi tiết, danh sách hiện có"],
        ["Booking", "/api/v1/bookings/mine, /api/v1/bookings/property", "Tạo booking và xem booking cá nhân"],
        ["Host", "/api/v1/host/properties, /profile-request, /bookings, /property-images", "Quản lý cơ sở và hồ sơ host"],
        ["Operator", "/api/v1/operator/*", "Dashboard, listings, bookings, host approvals, tasks, disputes"],
        ["Admin", "/api/v1/admin/*", "Dashboard, users, listings, bookings, payments, promotions, commissions, reviews, audit logs, operators"],
        ["Province", "/api/v1/provinces", "Danh mục tỉnh/thành và gán operator"],
        ["Notifications", "/api/v1/notifications/mine", "Thông báo cho user đang đăng nhập"],
    ],
    [Inches(1.0), Inches(3.8), Inches(2.7)],
)

add_heading(doc, "6. Database phản ánh đúng hiện trạng", 1)
add_table(
    doc,
    ["Bảng", "Vai trò"],
    [
        ["User", "Tài khoản hệ thống, role, profile, trạng thái, quan hệ operator/host/reviewer/assignee"],
        ["Province", "Tỉnh/thành phố, code, type, gán operator"],
        ["Property", "Cơ sở lưu trú với toàn bộ thông tin hiển thị, legal, pricing, rules, availability"],
        ["PropertyOwner / PropertyBedroom / PropertyImage / PropertyLanguage / PropertyRatePlan / PropertyChildPricing", "Chi tiết pháp lý, phòng, ảnh, ngôn ngữ, gói giá, giá trẻ em"],
        ["Tour / TourImage / TourItineraryDay / TourInclusion / TourAvailability", "Mô hình tour đã có trong schema"],
        ["Booking / Payment / Review", "Đặt chỗ, thanh toán, đánh giá"],
        ["Promotion / PromotionRedemption / CommissionRule", "Khuyến mãi, hoa hồng"],
        ["OperatorProvinceAssignment / OperatorTask / HostApprovalRequest / Dispute / AuditLog / Notification", "Lõi vận hành và quản trị"],
    ],
    [Inches(2.4), Inches(4.1)],
)

add_heading(doc, "7. Những phần cần ghi rõ là chưa hoàn thiện", 1)
for item in [
    "OAuth Google/Facebook: chỉ nên ghi là định hướng hoặc placeholder.",
    "2FA/OTP email: chưa hoàn thiện đầy đủ.",
    "Payment gateway VNPay/MoMo/QR: có model dữ liệu nhưng chưa phải flow thanh toán hoàn chỉnh.",
    "Xuất Excel/PDF: chưa có chức năng thực sự.",
    "Google Maps API: chưa tích hợp; form đăng cơ sở dùng OpenStreetMap/Nominatim.",
    "AES-256/Redis/Cloudinary: chưa có tích hợp đúng nghĩa trong code hiện tại.",
]:
    add_bullet(doc, item)

add_heading(doc, "8. Bảo mật và tuân thủ", 1)
add_table(
    doc,
    ["Trạng thái", "Mô tả"],
    [
        ["Đã có", "JWT, role-based access, bcrypt, validation đầu vào, audit log, rate limit ở một số luồng"],
        ["Cần chỉnh tài liệu", "Không nên ghi 2FA/OAuth/PCI DSS như đã hoàn thiện nếu code chưa có"],
        ["Tuân thủ", "Có thể mô tả theo hướng chuẩn bị hỗ trợ bảo vệ dữ liệu cá nhân, nhưng tránh khẳng định đã đạt đầy đủ nếu chưa có bằng chứng triển khai"],
    ],
    [Inches(1.4), Inches(5.1)],
)

add_heading(doc, "9. Đề xuất chỉnh trực tiếp các mục trong tài liệu cũ", 1)
for item in [
    "Chương 1-2: cập nhật câu mô tả tổng quan theo 3 app và hệ thống hiện tại, thêm Tour và quản trị operator.",
    "Chương 4: sửa actor thành 5 role thực tế.",
    "Chương 5: cập nhật chức năng theo module đã có và bỏ các claim chưa triển khai.",
    "Chương 6: sửa toàn bộ danh sách endpoint theo route thật.",
    "Chương 7: viết lại flow host approval, booking, admin/operator workflow theo code hiện tại.",
    "Chương 10: thay sơ bộ bằng mô hình database thực tế, không dùng bảng dự kiến cũ.",
    "Chương 11: hạ các mục OAuth/2FA/payment/export xuống trạng thái roadmap.",
]:
    add_number(doc, item)

add_heading(doc, "10. Lịch sử phiên bản", 1)
add_table(
    doc,
    ["Phiên bản", "Ngày", "Người viết", "Mô tả"],
    [
        ["v1.0", "05/2026", "Nhóm 4", "Bản mô tả ban đầu"],
        ["v1.1", "10/06/2026", "Codex cập nhật theo source hiện tại", "Đồng bộ đặc tả với hệ thống TripNest đã chỉnh sửa và bổ sung"],
    ],
    [Inches(0.8), Inches(1.0), Inches(2.2), Inches(2.5)],
)

doc.save(OUT)
print(OUT)
